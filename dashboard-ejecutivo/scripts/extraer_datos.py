"""Extrae del material fuente las cifras del Dashboard Ejecutivo MKT SOC 2026.

Uso:
  python3 -I extraer_datos.py <ruta_openpyxl> <salida.json> \
      --contable <Ejercicio_de_presupuesto.xlsx> --franquicias <...> --socleads <...> \
      --store <...> --precalificador <...>

Sólo lee los Excel (no los modifica). Toma únicamente cifras agregadas: no copia
datos personales (nombres, teléfonos, correos) al dashboard.
"""
import argparse
import json
import sys
from datetime import datetime

ap = argparse.ArgumentParser()
ap.add_argument('pylib')
ap.add_argument('salida')
ap.add_argument('--contable', required=True)
ap.add_argument('--franquicias', required=True)
ap.add_argument('--socleads', required=True)
ap.add_argument('--store', required=True)
ap.add_argument('--precalificador', required=True)
a = ap.parse_args()
sys.path.insert(0, a.pylib)
import openpyxl  # noqa: E402


def libro(ruta):
    return openpyxl.load_workbook(ruta, read_only=True, data_only=True)


def filas(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def r2(x):
    return round(float(x or 0), 2)


nombre = lambda ruta: ruta.replace('\\', '/').split('/')[-1].split('-', 1)[-1] if '-' in ruta.split('/')[-1][:9] else ruta.split('/')[-1]

# ───────────── Contable (fuente oficial) ─────────────
wb = libro(a.contable)
res = filas(wb['RESUMEN'])
hdr_i = next(i for i, r in enumerate(res) if r and r[0] == 'REPORTE MKT')
meses_txt = [m for m in res[hdr_i][1:9]]
categorias = []
total_fila = None
for r in res[hdr_i + 1:]:
    if not r or r[0] is None:
        continue
    if str(r[0]).strip() == '2026':
        total_fila = r
        break
    mensual = [r2(v) for v in r[1:9]]
    categorias.append({'nombre': str(r[0]).strip(), 'mensual': mensual, 'total': r2(r[9])})
assert total_fila is not None
total_contable = r2(total_fila[9])
suma_cat = round(sum(c['total'] for c in categorias), 2)
for c in categorias:
    assert abs(sum(c['mensual']) - c['total']) < 0.01, c['nombre']

vs = filas(wb['vs 2025'])
ingreso = next(r for r in vs if r and len(r) > 1 and r[1] == 'Ingreso Tienda')
ptto_total = next(r for r in vs if r and len(r) > 16 and r[1] is None and r[2] is not None and isinstance(r[14], (int, float)) and abs(r[14] - total_contable) < 0.01)
ptto_excel = {}
for r in vs:
    if r and len(r) > 16 and isinstance(r[1], str) and isinstance(r[16], (int, float)):
        ptto_excel[r[1].strip()] = r2(r[16])

g = wb['GASTOS']
gastos_total = 0.0
gastos_n = 0
gastos_cat = {}
hdr = None
for r in g.iter_rows(values_only=True):
    if hdr is None:
        if r and r[0] == 'SISTEMA':
            hdr = list(r)
            i_precio = hdr.index('PRECIO')
            i_rep = hdr.index('REPORTE MKT')
        continue
    if r[i_precio] is None:
        continue
    gastos_total += float(r[i_precio])
    gastos_n += 1
    k = str(r[i_rep]).strip()
    gastos_cat[k] = gastos_cat.get(k, 0.0) + float(r[i_precio])
contable = {
    'fuente': nombre(a.contable),
    'hoja': 'RESUMEN',
    'corte': '2026-08',
    'meses': ['2026-%02d' % m for m in range(1, 9)],
    'mesesTexto': meses_txt,
    'categorias': categorias,
    'total': total_contable,
    'sumaCategorias': suma_cat,
    'detalle': {'hoja': 'GASTOS', 'registros': gastos_n, 'total': r2(gastos_total), 'porCategoria': {k: r2(v) for k, v in gastos_cat.items()}},
    'ingresoTienda': {'mensual': [r2(v) for v in ingreso[2:10]], 'total': r2(ingreso[14])},
    'pttoExcelAgosto': {'total': r2(ptto_total[16]), 'porCategoria': ptto_excel},
}

# ───────────── Franquicias ─────────────
wb = libro(a.franquicias)
rt = filas(wb['Resumen trimestral'])
def fila_metrica(nombre_m):
    r = next(x for x in rt if x and x[0] == nombre_m)
    return [r2(v) for v in r[1:6]]
inv = fila_metrica('Inversión ($)')
leads = fila_metrica('Leads')
mql = fila_metrica('MQLs / Leads calificados')
sql = fila_metrica('SQLs / Tours informativos')
cierres = fila_metrica('Cierres')
em = filas(wb['Embudo mensual'])
mensual_fr = []
for r in em:
    if r and r[0] in ('Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio', 'Julio', 'Agosto', 'Septiembre'):
        mensual_fr.append({'mes': r[0], 'inversion': r2(r[2]), 'leads': int(r[5]), 'mql': int(r[6]), 'sql': int(r[7]), 'cierres': int(r[8])})
franquicias = {
    'fuente': nombre(a.franquicias),
    'periodo': 'Enero–septiembre 2026 (T1–T3)',
    'trimestres': ['T1', 'T2', 'T3'],
    'inversion': inv[:3], 'leads': leads[:3], 'mql': mql[:3], 'sql': sql[:3], 'cierres': cierres[:3],
    'total': {'inversion': inv[4], 'leads': leads[4], 'mql': mql[4], 'sql': sql[4], 'cierres': cierres[4]},
    'mensual': mensual_fr,
}

# ───────────── SOC Leads+ ─────────────
wb = libro(a.socleads)
rs = filas(wb['Resumen'])
def val(etq):
    return next(r[1] for r in rs if r and r[0] == etq)
embudo = []
for r in rs[4:15]:
    if r and r[3] and r[3] != 'TOTAL' and r[3] != 'Etapa':
        embudo.append({'etapa': r[3], 'leads': int(r[4])})
anios = {}
for r in wb['Detalle'].iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    y = r[1].year if isinstance(r[1], datetime) else None
    anios[str(y)] = anios.get(str(y), 0) + 1
socleads = {
    'fuente': nombre(a.socleads),
    'corte': '2026-10-16',
    'periodo': 'Negocios del programa al 16-oct-2026 (exportación HubSpot)',
    'oficinas': int(val('Oficinas en el programa')),
    'leads': int(val('Leads totales')),
    'sinGestionar': int(val('Sin gestionar (Asignado)')),
    'contactados': int(val('Contactados')),
    'enProceso': int(val('En proceso (Asesorado + Integración + Banco)')),
    'perdidos': int(val('Perdidos (No contactable, No responde, No interesado, No perfila, Descartado)')),
    'embudo': embudo,
    'registrosPorAnio': anios,
}

# ───────────── SOC Store (reporte comercial) ─────────────
wb = libro(a.store)
ini = filas(wb['Inicio'])
periodo_store = next(r[1] for r in ini if r and len(r) > 1 and r[1] and 'Reporte de estadísticas' in str(r[1]))
por_fecha = []
for r in filas(wb['Por fecha'])[4:]:
    if r and r[0]:
        por_fecha.append({'mes': r[0], 'ventas': int(r[1] or 0), 'unidades': int(r[2] or 0), 'facturacion': r2(r[3])})
pago = filas(wb['Pago y financiación'])
desc100 = next(r for r in pago if r and r[0] == 'Pedido con 100% de descuento')
cc = next(r for r in pago if r and r[0] == 'CENTRO DE COSTO')
resumen_store = {str(r[1]).strip(): str(r[0]).strip() for r in filas(wb['Resumen']) if r and r[0] and r[1]}
def num(txt):
    return float(str(txt).replace('$', '').replace('.', '').replace(',', '.').strip())
store = {
    'fuente': nombre(a.store),
    'periodo': periodo_store.replace('Reporte de estadísticas del ', ''),
    'transacciones': int(num(resumen_store['ventas'])),
    'facturado': num(resumen_store['total facturado']),
    'unidades': int(num(resumen_store['productos vendidos'])),
    'ticketPromedio': num(resumen_store['por venta (ticket promedio)']),
    'ticketMediano': num(resumen_store['ticket mediano']),
    'pedidosDescuentoTotal': int(desc100[1]),
    'ventasCentroCosto': {'ventas': int(cc[1]), 'facturacion': r2(cc[3])},
    'porMes': por_fecha,
}

# ───────────── Precalificador ─────────────
wb = libro(a.precalificador)
par = filas(wb['Parámetros'])
pv = lambda etq: next(r[1] for r in par if r and r[0] == etq)
zm = filas(wb['Zona x Mes'])
cab = zm[0]
tot = next(r for r in zm if r and r[0] == 'Total nacional')
pend = next(r for r in zm if r and r[0] == 'Zona pendiente')
meses_p = [c for c in cab[1:23]]
precal = {
    'fuente': nombre(a.precalificador),
    'corte': '2026-10-15',
    'historico': int(pv("Registros en 'Datos homologados'")),
    'sinZona': int(pv('Registros en Zona pendiente')),
    'meses': meses_p,
    'mensual': [int(v) for v in tot[1:23]],
    'sinZonaMensual': [int(v) for v in pend[1:23]],
    'total2025': int(tot[cab.index('Total 2025')]),
    'total2026': int(tot[cab.index('Total 2026')]),
    'sinZona2026': int(pend[cab.index('Total 2026')]),
}
assert precal['total2025'] + precal['total2026'] == precal['historico']

out = {
    'generado': datetime.now().strftime('%Y-%m-%d %H:%M'),
    'contable': contable,
    'franquicias': franquicias,
    'socleads': socleads,
    'store': store,
    'precalificador': precal,
}
with open(a.salida, 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=1)

print('Contable total', total_contable, 'suma categorías', suma_cat, 'GASTOS', r2(gastos_total), gastos_n, 'registros')
print('PTTO Excel a agosto', contable['pttoExcelAgosto']['total'])
for k, v in sorted(gastos_cat.items()):
    c = next((x for x in categorias if x['nombre'] == k), None)
    print('  GASTOS vs RESUMEN', k, r2(v), c and c['total'])
print('Franquicias', franquicias['total'])
print('SOCLeads', {k: socleads[k] for k in ('oficinas', 'leads', 'sinGestionar', 'contactados', 'enProceso', 'perdidos')}, anios)
print('Store', {k: store[k] for k in ('periodo', 'transacciones', 'facturado', 'unidades', 'ticketPromedio', 'ticketMediano', 'pedidosDescuentoTotal')})
print('Precal', {k: precal[k] for k in ('historico', 'sinZona', 'total2025', 'total2026', 'sinZona2026')})
